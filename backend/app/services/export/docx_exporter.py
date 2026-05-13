from io import BytesIO
from typing import Dict, Optional
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.enum.section import WD_ORIENT
from .asset_manager import AssetManager
import re

class DOCXExporter:
    def __init__(self):
        self.asset_manager = AssetManager()
    
    def _add_heading(self, doc, text, level):
        heading = doc.add_heading(text, level=level)
        heading.runs[0].font.size = Pt(12 + (4 - level) * 2)
        heading.runs[0].font.bold = True
        heading.paragraph_format.space_after = Pt(6)
    
    def _add_paragraph(self, doc, text, style=None):
        p = doc.add_paragraph(text)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = Pt(18)
        if style == "code":
            p.style = "No Spacing"
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0, 0, 128)
        return p
    
    def _add_image(self, doc, image_path, doc_id):
        if image_path.startswith("data:image/"):
            import base64
            data = image_path.split(",")[1]
            img_bytes = base64.b64decode(data)
            img_stream = BytesIO(img_bytes)
            doc.add_picture(img_stream, width=Inches(5))
        else:
            if image_path.startswith("/"):
                full_path = self.asset_manager.static_dir / image_path.lstrip("/")
            else:
                full_path = self.asset_manager.static_dir / "images" / doc_id / image_path
            
            if full_path.exists():
                doc.add_picture(str(full_path), width=Inches(5))
    
    def _parse_content(self, content, doc, doc_id):
        lines = content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            
            if line.startswith("$$") and line.endswith("$$"):
                math_content = line[2:-2].strip()
                p = doc.add_paragraph(math_content)
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(12)
                i += 1
                continue
            
            if line.startswith("$") and line.endswith("$"):
                p = doc.add_paragraph(line)
                i += 1
                continue
            
            if line.startswith("# "):
                self._add_heading(doc, line[2:], level=1)
                i += 1
                continue
            
            if line.startswith("## "):
                self._add_heading(doc, line[3:], level=2)
                i += 1
                continue
            
            if line.startswith("### "):
                self._add_heading(doc, line[4:], level=3)
                i += 1
                continue
            
            if line.startswith("```"):
                code_block = []
                i += 1
                while i < len(lines) and not lines[i].startswith("```"):
                    code_block.append(lines[i])
                    i += 1
                i += 1
                code_text = "\n".join(code_block)
                p = doc.add_paragraph(code_text)
                p.style = "No Spacing"
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(10)
                continue
            
            if line.startswith("> "):
                p = doc.add_paragraph(line[2:])
                p.paragraph_format.left_indent = Inches(0.5)
                p.style = "Quote"
                i += 1
                continue
            
            img_match = re.search(r'!\[([^\]]*)\]\(([^)]+)\)', line)
            if img_match:
                self._add_image(doc, img_match.group(2), doc_id)
                i += 1
                continue
            
            if line:
                self._add_paragraph(doc, line)
            i += 1
    
    async def export(self, original_text: str, translated_text: str, doc_id: str, title: str = "Translation",
                    include_original: bool = True, include_translation: bool = True) -> Dict[str, bytes]:
        
        doc = Document()
        
        section = doc.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        
        if title:
            title_para = doc.add_paragraph(title)
            title_para.runs[0].font.size = Pt(16)
            title_para.runs[0].font.bold = True
            title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            title_para.paragraph_format.space_after = Pt(12)
        
        if include_original:
            self._add_heading(doc, "Original", level=1)
            self._parse_content(original_text, doc, doc_id)
        
        if include_translation:
            if include_original:
                doc.add_page_break()
            self._add_heading(doc, "Translation", level=1)
            self._parse_content(translated_text, doc, doc_id)
        
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return {
            "content": buffer.read(),
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "extension": ".docx"
        }
