from __future__ import annotations
from typing import List, Optional
from io import BytesIO
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, Emu
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.models.block_schema import (
    Block, BlockType, InlineNode, InlineType,
    Document, TableAlign
)
from pathlib import Path
import re
import base64


class DocxRenderer:

    def __init__(self, asset_base_path: Optional[Path] = None):
        self.asset_base_path = asset_base_path

    def render(self, document: Document, title: str = "") -> bytes:
        doc = DocxDocument()

        section = doc.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)

        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)

        if title:
            title_para = doc.add_paragraph()
            title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = title_para.add_run(title)
            run.font.size = Pt(16)
            run.font.bold = True
            title_para.paragraph_format.space_after = Pt(12)

        for block in document.blocks:
            self._render_block_to_docx(doc, block)

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.read()

    def _render_block_to_docx(self, doc: DocxDocument, block: Block):
        if block.type == BlockType.heading:
            heading = doc.add_heading('', level=min(block.level, 3))
            self._add_inlines_to_paragraph(heading, block.inlines, doc)
            size_map = {1: 20, 2: 18, 3: 16}
            pt = size_map.get(block.level, 14)
            for run in heading.runs:
                run.font.size = Pt(pt)
                run.font.bold = True
            heading.paragraph_format.space_after = Pt(6)

        elif block.type == BlockType.paragraph:
            text = self._inlines_to_plain_text(block.inlines)
            if not text.strip():
                return
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = Pt(18)
            self._add_inlines_to_paragraph(p, block.inlines, doc)

        elif block.type == BlockType.code_block:
            p = doc.add_paragraph()
            p.style = doc.styles['No Spacing']
            for line in block.content.split('\n'):
                if line:
                    run = p.add_run(line + '\n')
                    run.font.name = 'Consolas'
                    run.font.size = Pt(10)

        elif block.type == BlockType.math_block:
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(f'[ {block.content} ]')
            run.font.size = Pt(12)
            run.font.italic = True

        elif block.type == BlockType.bullet_list:
            for child in block.children:
                p = doc.add_paragraph(style='List Bullet')
                self._add_inlines_to_paragraph(p, child.inlines, doc)

        elif block.type == BlockType.ordered_list:
            for idx, child in enumerate(block.children):
                p = doc.add_paragraph(style='List Number')
                self._add_inlines_to_paragraph(p, child.inlines, doc)

        elif block.type == BlockType.blockquote:
            for child in block.children:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.style = doc.styles['Normal']
                self._add_inlines_to_paragraph(p, child.inlines, doc)

        elif block.type == BlockType.table:
            self._render_table_to_docx(doc, block)

        elif block.type == BlockType.thematic_break:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('_' * 60)
            run.font.size = Pt(8)
            run.font.color.rgb = None

        elif block.type == BlockType.html_block:
            pass

    def _render_table_to_docx(self, doc: DocxDocument, block: Block):
        if not block.rows or not block.rows[0]:
            return

        rows_data = block.rows
        num_cols = max(len(row) for row in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for ri, row in enumerate(rows_data):
            for ci, cell_inlines in enumerate(row):
                if ci >= num_cols:
                    break
                cell = table.rows[ri].cells[ci]
                cell.text = ''
                p = cell.paragraphs[0]
                self._add_inlines_to_paragraph(p, cell_inlines, doc)

    def _add_inlines_to_paragraph(self, paragraph, inlines: List[InlineNode], doc: Optional[DocxDocument] = None):
        for node in inlines:
            self._add_inline_to_paragraph(paragraph, node, doc)

    def _add_inline_to_paragraph(self, paragraph, node: InlineNode, doc: Optional[DocxDocument] = None):
        if node.type == InlineType.text:
            run = paragraph.add_run(node.content)
            run.font.size = Pt(12)

        elif node.type == InlineType.bold:
            inner = self._render_inlines_text(node) if node.children else node.content
            run = paragraph.add_run(inner)
            run.bold = True
            run.font.size = Pt(12)

        elif node.type == InlineType.italic:
            inner = self._render_inlines_text(node) if node.children else node.content
            run = paragraph.add_run(inner)
            run.italic = True
            run.font.size = Pt(12)

        elif node.type == InlineType.underline:
            run = paragraph.add_run(node.content)
            run.underline = True
            run.font.size = Pt(12)

        elif node.type == InlineType.code:
            run = paragraph.add_run(node.content)
            run.font.name = 'Consolas'
            run.font.size = Pt(10)

        elif node.type == InlineType.math:
            run = paragraph.add_run(f'({node.content})')
            run.font.size = Pt(12)
            run.font.italic = True

        elif node.type == InlineType.link:
            inner = self._render_inlines_text(node) if node.children else node.content
            run = paragraph.add_run(f'{inner} ({node.url})')
            run.font.size = Pt(12)

        elif node.type == InlineType.image:
            if doc:
                self._add_image_to_docx(doc, node.url)

    def _add_image_to_docx(self, doc: DocxDocument, src: str):
        try:
            if src.startswith('data:image/') or src.startswith('data:'):
                data = src.split(',')[1]
                img_bytes = base64.b64decode(data)
                img_stream = BytesIO(img_bytes)
                doc.add_picture(img_stream, width=Inches(4))
            elif self.asset_base_path:
                if src.startswith('/'):
                    full_path = self.asset_base_path / src.lstrip('/')
                else:
                    full_path = self.asset_base_path / src
                if full_path.exists():
                    doc.add_picture(str(full_path), width=Inches(4))
        except Exception:
            pass

    def _render_inlines_text(self, node: InlineNode) -> str:
        if node.type == InlineType.text:
            return node.content
        parts: List[str] = []
        for child in node.children:
            parts.append(self._render_inlines_text(child))
        return ''.join(parts)

    def _inlines_to_plain_text(self, inlines: List[InlineNode]) -> str:
        parts: List[str] = []
        for node in inlines:
            parts.append(self._render_inlines_text(node))
        return ''.join(parts)